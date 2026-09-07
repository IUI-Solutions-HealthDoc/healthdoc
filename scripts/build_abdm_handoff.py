#!/usr/bin/env python3
"""Render the reviewed ABDM handoff Markdown as an editable Word document.

Requires python-docx in a document-tooling environment, not the application.
Usage: python build_abdm_handoff.py [source.md] [output.docx]
Only headings, paragraphs, lists, tables, code fences and page breaks are used.
No Markdown content is executed and no network resources are fetched.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT / "docs/ABDM-M1-M2-M3-Execution-Requirements-2026-09-07.md"
NAVY = "092A52"
TEAL = "087D83"


def shade(element, color):
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), color)
    element.append(item)


def inline(paragraph, value):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", value):
        if not part:
            continue
        run = paragraph.add_run(part.strip("`") if part.startswith("`") else part)
        if part.startswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("`"):
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)


def make_table(document, lines):
    values = [
        [cell.strip() for cell in row.strip().strip("|").split("|")] for row in lines
    ]
    values = [row for row in values if not all(re.fullmatch(r":?-+:?", c) for c in row)]
    table = document.add_table(rows=0, cols=len(values[0]))
    table.autofit = False
    proportions = [0.36, 0.64] if len(values[0]) == 2 else [0.30, 0.23, 0.47]
    for col, width in zip(table.columns, proportions):
        col.width = Inches(6.85 * width)
    for index, row in enumerate(values):
        cells = table.add_row().cells
        properties = table.rows[-1]._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for cell, text, width in zip(cells, row, proportions):
            cell.width = Inches(6.85 * width)
            shade(
                cell._tc.get_or_add_tcPr(),
                NAVY if index == 0 else ("EEF5F8" if index % 2 else "F8FAFC"),
            )
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(10)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(source, output):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.71)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.27)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("243449")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.03
    for name, size, color in [
        ("Title", 25, NAVY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, TEAL),
        ("Heading 3", 11.5, TEAL),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "HEALTHDOC   /   ABDM SANDBOX EXECUTION HANDOFF"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.text = "7 September 2026  •  Engineering readiness ≠ NHA acceptance                       Page "
    footer.runs[0].font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = (
        "HealthDoc ABDM M1, M2 and M3 — Execution and Requirements"
    )
    document.core_properties.subject = (
        "Verified readiness, blockers, operator steps and exit evidence"
    )
    document.core_properties.author = "HealthDoc"
    document.core_properties.keywords = "ABDM, M1, M2, M3, HIP, HIU, sandbox"

    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    first_title = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            document.add_page_break()
        elif line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Inches(0.08)
            shade(p._p.get_or_add_pPr(), "F0F4F7")
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            make_table(document, table_lines)
            continue
        elif line.startswith("#"):
            depth = len(line) - len(line.lstrip("#"))
            style = "Title" if first_title else f"Heading {min(depth, 3)}"
            inline(document.add_paragraph(style=style), line[depth:].strip())
            first_title = False
        else:
            p = document.add_paragraph(
                style="List Bullet" if line.startswith("- ") else "Normal"
            )
            if line.startswith("- "):
                line = line[2:]
            elif re.match(r"^\d+\. ", line):
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.first_line_indent = Inches(-0.18)
            if line.startswith("https://"):
                # A clickable reference, without fetching the destination.
                p.paragraph_format.space_after = Pt(9)
                link = OxmlElement("w:hyperlink")
                link.set(
                    qn("r:id"), p.part.relate_to(line, RT.HYPERLINK, is_external=True)
                )
                run = OxmlElement("w:r")
                properties = OxmlElement("w:rPr")
                size = OxmlElement("w:sz")
                size.set(qn("w:val"), "17")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), TEAL)
                properties.extend([size, color])
                text = OxmlElement("w:t")
                text.text = line
                run.extend([properties, text])
                link.append(run)
                p._p.append(link)
            else:
                inline(p, line)
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Wrote {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, nargs="?", default=DEFAULT)
    parser.add_argument("output", type=Path, nargs="?")
    args = parser.parse_args()
    build(args.source, args.output or args.source.with_suffix(".docx"))
